#!/usr/bin/env python3
"""Chuyển tài liệu SRS (Markdown) sang Word (.docx).

Hỗ trợ: tiêu đề #/##/###, bảng, gạch đầu dòng, đoạn văn, in đậm **...**,
mã inline `...`, khối code ```...``` và tô đậm dòng tiêu đề của bảng.
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

MD_PATH = Path(__file__).parent / "SRS_SmartSpend.md"
OUT_PATH = Path(__file__).parent / "SRS_SmartSpend.docx"

BOLD_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")
BRAND = RGBColor(0x0D, 0x94, 0x88)


def add_inline(paragraph, text):
    """Thêm text có xử lý **đậm** và `mã`."""
    for part in BOLD_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x60)
        else:
            paragraph.add_run(part)


def is_separator_row(line):
    return set(line.replace("|", "").replace("-", "").replace(":", "").replace(" ", "")) == set()


def flush_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.paragraphs[0].text = ""
            text = row[j] if j < len(row) else ""
            add_inline(cell.paragraphs[0], text)
            if i == 0:  # header row -> bold
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()


def flush_code(doc, lines):
    p = doc.add_paragraph()
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.left_indent = Pt(6)


def main():
    content = MD_PATH.read_text(encoding="utf-8")
    doc = Document()

    # Font mặc định
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    table_rows = []
    in_table = False
    in_code = False
    code_lines = []
    title_done = False

    def flush_pending():
        nonlocal in_table, table_rows
        if in_table and table_rows:
            flush_table(doc, table_rows)
            table_rows = []
            in_table = False

    for raw in content.splitlines():
        line = raw.rstrip()

        # Khối code
        if line.strip().startswith("```"):
            if in_code:
                flush_code(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                flush_pending()
                in_code = True
            continue
        if in_code:
            code_lines.append(raw)
            continue

        if not line or line == "---":
            flush_pending()
            continue

        if line.startswith("# "):
            flush_pending()
            text = line[2:]
            if not title_done:
                h = doc.add_heading(text, level=0)
                h.alignment = WD_ALIGN_PARAGRAPH.CENTER
                title_done = True
            else:
                doc.add_heading(text, level=1)
        elif line.startswith("## "):
            flush_pending()
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            flush_pending()
            doc.add_heading(line[4:], level=2)
        elif line.startswith("|") and "|" in line[1:]:
            if is_separator_row(line):
                continue
            in_table = True
            cells = [c.strip() for c in line.strip("|").split("|")]
            table_rows.append(cells)
        elif line.startswith("- ") or line.startswith("* "):
            flush_pending()
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            flush_pending()
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\.\s", "", line))
        elif line.startswith("> "):
            flush_pending()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            run_prefix = p.add_run("")
            add_inline(p, line[2:])
            for r in p.runs:
                r.italic = True
        else:
            flush_pending()
            p = doc.add_paragraph()
            add_inline(p, line)

    flush_pending()
    if in_code and code_lines:
        flush_code(doc, code_lines)

    doc.save(OUT_PATH)
    print(f"Created: {OUT_PATH}")


if __name__ == "__main__":
    main()
