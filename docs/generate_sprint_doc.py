#!/usr/bin/env python3
"""Generate Sprint Backlog Word document from markdown source."""

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    import subprocess
    import sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

from pathlib import Path

MD_PATH = Path(__file__).parent / "Sprint_Backlog_SmartSpend.md"
OUT_PATH = Path(__file__).parent / "Sprint_Backlog_SmartSpend.docx"


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def main():
    content = MD_PATH.read_text(encoding="utf-8")
    doc = Document()

    # Title
    title = doc.add_heading("SMARTSPEND — SPRINT BACKLOG & SUB-TASK", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("Dự án tốt nghiệp | Bắt đầu: 01/06/2026 | 1 Sprint = 1 tuần")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    in_table = False
    table_rows = []

    for line in content.splitlines():
        line = line.rstrip()
        if not line or line == "---":
            if in_table and table_rows:
                _flush_table(doc, table_rows)
                table_rows = []
                in_table = False
            continue

        if line.startswith("# "):
            if in_table and table_rows:
                _flush_table(doc, table_rows)
                table_rows = []
                in_table = False
            add_heading(doc, line[2:], 1)
        elif line.startswith("## "):
            if in_table and table_rows:
                _flush_table(doc, table_rows)
                table_rows = []
                in_table = False
            add_heading(doc, line[3:], 2)
        elif line.startswith("### "):
            if in_table and table_rows:
                _flush_table(doc, table_rows)
                table_rows = []
                in_table = False
            add_heading(doc, line[4:], 3)
        elif line.startswith("|") and "|" in line[1:]:
            if set(line.replace("|", "").replace("-", "").replace(" ", "")) == set():
                continue
            in_table = True
            cells = [c.strip() for c in line.strip("|").split("|")]
            table_rows.append(cells)
        elif line.startswith("- "):
            if in_table and table_rows:
                _flush_table(doc, table_rows)
                table_rows = []
                in_table = False
            text = line[2:]
            add_bullet(doc, text)
        else:
            if in_table and table_rows:
                _flush_table(doc, table_rows)
                table_rows = []
                in_table = False
            if line.startswith("**") and line.endswith("**"):
                p = doc.add_paragraph()
                run = p.add_run(line.strip("*"))
                run.bold = True
            else:
                doc.add_paragraph(line)

    if in_table and table_rows:
        _flush_table(doc, table_rows)

    doc.save(OUT_PATH)
    print(f"Created: {OUT_PATH}")


def _flush_table(doc, rows):
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            if j < cols:
                table.rows[i].cells[j].text = cell
    doc.add_paragraph()


if __name__ == "__main__":
    main()
